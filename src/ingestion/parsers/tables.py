"""Tabular parsing: .csv (stdlib) and .xlsx (openpyxl).

Each data row is rendered as ``Header: value | ...`` and rows are packed into
pseudo-pages of ``TABLE_ROWS_PER_PAGE`` so every citation points at a small,
readable slice of the sheet.
"""
from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from langchain_core.documents import Document

from config import settings
from config.settings import MAX_TABLE_ROWS, TABLE_ROWS_PER_PAGE

from .common import (
    make_document,
    render_table_lines,
    read_text_bytes,
    validate_file,
)

CSV_EXTENSIONS = (".csv",)
XLSX_EXTENSIONS = (".xlsx",)

_XLSX_ROOT_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_XLSX_MERGE_NS = {"a": _XLSX_ROOT_NS}


def _sheet_merge_ranges(workbook, index: int) -> list[tuple[int, int, int, int]]:
    """Return ``(min_col, min_row, max_col, max_row)`` for every merged range of
    the worksheet at ``index``.

    ``read_only`` worksheets do not expose ``ws.merged_cells``, so the ranges are
    read straight from the sheet XML already held in the workbook archive.
    Coordinates are 1-based, matching ``iter_rows`` output; any sheet that fails
    to parse simply contributes no ranges.
    """
    archive = getattr(workbook, "_archive", None)
    if archive is None:
        return []

    try:
        root = ET.fromstring(
            archive.read(f"xl/worksheets/sheet{index + 1}.xml")
        )
    except (KeyError, OSError, ET.ParseError):
        return []

    merges = root.find("a:mergeCells", _XLSX_MERGE_NS)
    if merges is None:
        return []

    from openpyxl.utils import column_index_from_string

    ranges: list[tuple[int, int, int, int]] = []
    for node in merges.findall("a:mergeCell", _XLSX_MERGE_NS):
        ref = (node.get("ref") or "").strip()
        if ":" not in ref:
            continue
        start_ref, end_ref = ref.split(":", 1)
        start_match = re.match(r"([A-Z]+)(\d+)", start_ref.upper())
        end_match = re.match(r"([A-Z]+)(\d+)", end_ref.upper())
        if not start_match or not end_match:
            continue
        try:
            start_col = column_index_from_string(start_match.group(1))
            end_col = column_index_from_string(end_match.group(1))
        except ValueError:
            continue
        ranges.append(
            (start_col, int(start_match.group(2)), end_col, int(end_match.group(2)))
        )
    return ranges


def _fill_merged_cells(
    rows: list[list[str]],
    ranges: list[tuple[int, int, int, int]],
) -> list[list[str]]:
    """Write each merged range's top-left value into every covered cell.

    OpenPyXL stores a merged range's value only in its top-left cell (the rest
    come back empty); repeating that value reproduces merged headers/category
    labels across the whole range.
    """
    for start_col, start_row, end_col, end_row in ranges:
        try:
            top_left = rows[start_row - 1][start_col - 1]
        except IndexError:
            continue
        for row_index in range(start_row, end_row + 1):
            if row_index - 1 >= len(rows):
                break
            row = rows[row_index - 1]
            for col_index in range(start_col, end_col + 1):
                if col_index - 1 < len(row):
                    row[col_index - 1] = top_left
    return rows


def _workbook_metadata(workbook) -> dict:
    """Workbook property metadata (title/creator/created) if present."""

    if not settings.EXTRACT_EMBEDDED_METADATA:
        return {}
    try:
        props = workbook.properties
    except Exception:
        return {}
    meta: dict = {}
    if props.title:
        meta["doc_title"] = str(props.title).strip()
    if props.creator:
        meta["doc_author"] = str(props.creator).strip()
    if props.created:
        meta["doc_created"] = props.created.isoformat()
    if props.modified:
        meta["doc_modified"] = props.modified.isoformat()
    return meta


def _rows_to_documents(
    rows: list[list[str]],
    *,
    filename: str,
    rows_per_page: int = TABLE_ROWS_PER_PAGE,
    sheet: str | None = None,
    page_start: int = 1,
    doc_meta: dict | None = None,
) -> list[Document]:
    lines = render_table_lines(rows)
    if not lines:
        return []

    if len(lines) > MAX_TABLE_ROWS:
        lines = lines[:MAX_TABLE_ROWS]

    prefix = f"[Sheet: {sheet}]\n" if sheet else ""
    pages = [
        prefix + "\n".join(lines[i : i + rows_per_page])
        for i in range(0, len(lines), rows_per_page)
    ]

    extra = {"sheet": sheet} if sheet else {}
    extra.update(doc_meta or {})
    return [
        make_document(
            page_text,
            filename=filename,
            page_number=page_number,
            raw_chars=len(page_text),
            clean_chars=len(page_text),
            **extra,
        )
        for page_number, page_text in enumerate(pages, start=page_start)
    ]


def parse_csv(
    file_path: str | Path,
    rows_per_page: int = TABLE_ROWS_PER_PAGE,
) -> list[Document]:
    """Parse a CSV file; the first row is treated as the header."""

    path = validate_file(file_path, CSV_EXTENSIONS)

    # Decode with the same fallback chain as other text formats (utf-8, utf-8-sig,
    # latin-1, replace) instead of assuming utf-8, so cp1252/latin-1 CSVs parse.
    text = read_text_bytes(path.read_bytes())
    rows = list(csv.reader(io.StringIO(text)))

    return _rows_to_documents(rows, filename=path.name, rows_per_page=rows_per_page)


def parse_xlsx(
    file_path: str | Path,
    rows_per_page: int = TABLE_ROWS_PER_PAGE,
) -> list[Document]:
    """Parse an .xlsx workbook; every sheet becomes its own set of pages."""

    path = validate_file(file_path, XLSX_EXTENSIONS)

    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    doc_meta = _workbook_metadata(workbook)

    try:
        documents: list[Document] = []
        # Page numbers run across the whole workbook: chunk ids and citation
        # keys are (document_id, page_number), so per-sheet restarts would
        # collide.
        page_start = 1
        for index, sheet in enumerate(workbook.worksheets):
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = _fill_merged_cells(rows, _sheet_merge_ranges(workbook, index))
            sheet_documents = _rows_to_documents(
                rows,
                filename=path.name,
                rows_per_page=rows_per_page,
                sheet=sheet.title,
                page_start=page_start,
                doc_meta=doc_meta,
            )
            documents.extend(sheet_documents)
            page_start += len(sheet_documents)
        return documents
    finally:
        workbook.close()
