"""PPTX parsing via python-pptx: one pseudo-page per slide."""
from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from config import settings

from .common import (
    clean_page_text,
    make_document,
    render_table_lines,
    validate_file,
)

ALLOWED_EXTENSIONS = (".pptx",)


def _embedded_metadata(presentation) -> dict:
    """Core-property metadata (title/author/created) if present."""

    if not settings.EXTRACT_EMBEDDED_METADATA:
        return {}
    try:
        props = presentation.core_properties
    except Exception:
        return {}
    meta: dict = {}
    if props.title:
        meta["doc_title"] = str(props.title).strip()
    if props.author:
        meta["doc_author"] = str(props.author).strip()
    if props.created:
        meta["doc_created"] = props.created.isoformat()
    if props.modified:
        meta["doc_modified"] = props.modified.isoformat()
    return meta


def _slide_text(slide) -> tuple[str, int, int]:
    """Return ``(text, raw_chars, clean_chars)`` for one slide."""

    parts: list[str] = []
    raw_total = clean_total = 0

    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                text, stats = clean_page_text(paragraph.text)
                if text:
                    parts.append(text)
                    raw_total += stats["raw_chars"]
                    clean_total += stats["clean_chars"]
        if getattr(shape, "has_table", False):
            rows = [
                [cell.text for cell in row.cells]
                for row in shape.table.rows
            ]
            for line in render_table_lines(rows):
                parts.append(line)
                raw_total += len(line)
                clean_total += len(line)

    if getattr(slide, "has_notes_slide", False):
        notes, stats = clean_page_text(slide.notes_slide.notes_text_frame.text)
        if notes:
            parts.append(f"Notes: {notes}")
            raw_total += stats["raw_chars"] + len("Notes: ")
            clean_total += stats["clean_chars"] + len("Notes: ")

    text, _ = clean_page_text("\n".join(parts))
    return text, raw_total, clean_total


def parse_pptx(file_path: str | Path) -> list[Document]:
    """Parse a .pptx deck; each non-empty slide is one page (slide number)."""

    path = validate_file(file_path, ALLOWED_EXTENSIONS)

    from pptx import Presentation

    presentation = Presentation(str(path))
    doc_meta = _embedded_metadata(presentation)

    documents: list[Document] = []
    empty = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text, raw_chars, clean_chars = _slide_text(slide)
        if not text:
            empty += 1
            continue
        documents.append(
            make_document(
                text,
                filename=path.name,
                page_number=slide_number,
                doc_type="pptx",
                raw_chars=raw_chars,
                clean_chars=clean_chars,
                **doc_meta,
            )
        )

    if not documents:
        raise ValueError(
            f"No text could be extracted from {path.name} "
            f"({empty} slide(s) contained no text)."
        )

    return documents
